"""Verify local document extraction for the file-attach feature.

The extractor is Qt-free and blocking — it runs on the chat window's worker
thread. These tests use real PDF bytes (hand-built via :mod:`pdf_fixtures`) and
real ``.docx`` files (built with python-docx) so they cover the actual parsing
paths, not mocks.

Covered here: single/multi-page PDF text, scanned (text-less) PDF detection,
unreadable PDFs, DOCX paragraphs/headings/lists/tables in document order,
unsupported formats, the large-document cap, and the source-grounded prompt
composer.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from aashi.ui.services.document_input import (
    DOCUMENT_EXTS,
    MAX_DOCUMENT_CHARS,
    DocumentInputError,
    ScannedDocumentError,
    UnsupportedDocumentError,
    build_document_prompt,
    cap_document_text,
    extract_document,
    extract_pdf_text,
    is_document_path,
)
from pdf_fixtures import make_pdf


# ── DOCX fixture builder ─────────────────────────────────────────────────────


def _build_docx(path: Path, blocks: list[tuple[str, object]]) -> Path:
    """Write ``path`` as a .docx composed of ``blocks``.

    Each block is ``(kind, payload)`` where kind is one of ``"p"`` (paragraph
    text), ``"h"`` (``(level, text)``), ``"bullet"`` (text), ``"number"``
    (text) or ``"table"`` (list-of-lists of cell text).
    """
    from docx import Document

    document = Document()
    for kind, payload in blocks:
        if kind == "p":
            document.add_paragraph(str(payload))
        elif kind == "h":
            level, text = payload  # type: ignore[misc]
            document.add_heading(str(text), level=level)
        elif kind == "bullet":
            document.add_paragraph(str(payload), style="List Bullet")
        elif kind == "number":
            document.add_paragraph(str(payload), style="List Number")
        elif kind == "table":
            rows = list(payload)  # type: ignore[arg-type]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    table.cell(r, c).text = str(cell)
    document.save(str(path))
    return path


# ── PDF extraction ───────────────────────────────────────────────────────────


def test_pdf_single_page_text_is_extracted(tmp_path) -> None:
    pdf = tmp_path / "notes.pdf"
    pdf.write_bytes(make_pdf(["Photosynthesis converts sunlight into energy."]))

    text = extract_pdf_text(str(pdf))

    assert "Photosynthesis converts sunlight into energy." in text


def test_pdf_multipage_text_is_extracted_in_order(tmp_path) -> None:
    pdf = tmp_path / "book.pdf"
    pdf.write_bytes(
        make_pdf(
            [
                "First page: the introduction.",
                "Second page: the main topic.",
                "Third page: the conclusion.",
            ]
        )
    )

    text = extract_pdf_text(str(pdf))

    # Every page's content survives, with its page marker.
    assert "[Page 1]" in text
    assert "[Page 2]" in text
    assert "[Page 3]" in text
    assert text.index("the introduction") < text.index("the main topic") < text.index(
        "the conclusion"
    )


def test_pdf_without_text_is_reported_as_scanned(tmp_path) -> None:
    pdf = tmp_path / "scan.pdf"
    pdf.write_bytes(make_pdf(["", ""]))

    with pytest.raises(ScannedDocumentError):
        extract_pdf_text(str(pdf))


def test_pdf_unreadable_raises_document_error(tmp_path) -> None:
    pdf = tmp_path / "broken.pdf"
    pdf.write_bytes(b"this is definitely not a pdf file")

    with pytest.raises(DocumentInputError):
        extract_pdf_text(str(pdf))


def test_pdf_page_failure_is_skipped_not_fatal(monkeypatch) -> None:
    """A page whose content stream is corrupt is skipped, not fatal."""
    from unittest.mock import MagicMock

    good_page = MagicMock()
    good_page.extract_text.return_value = "Real page text."
    broken_page = MagicMock()
    broken_page.extract_text.side_effect = RuntimeError("broken stream")

    fake_reader = MagicMock()
    fake_reader.pages = [good_page, broken_page]

    monkeypatch.setattr(
        "aashi.ui.services.document_input.PdfReader", lambda _path: fake_reader
    )

    text = extract_pdf_text("ignored.pdf")

    # The good page survives; the broken one is skipped without sinking the doc.
    assert "Real page text." in text
    assert "[Page 2]" not in text


# ── DOCX extraction ──────────────────────────────────────────────────────────


def test_docx_paragraphs_are_extracted_in_order(tmp_path) -> None:
    docx = _build_docx(
        tmp_path / "notes.docx",
        [("p", "First paragraph."), ("p", "Second paragraph.")],
    )

    text = extract_document(str(docx))

    assert text.index("First paragraph.") < text.index("Second paragraph.")


def test_docx_headings_are_marked_with_hashes(tmp_path) -> None:
    docx = _build_docx(
        tmp_path / "headings.docx",
        [
            ("h", (1, "Main Findings")),
            ("p", "Body text."),
            ("h", (2, "Subsection")),
        ],
    )

    text = extract_document(str(docx))

    assert "# Main Findings" in text
    assert "## Subsection" in text
    assert text.index("# Main Findings") < text.index("Body text.")


def test_docx_bullet_lists_are_marked(tmp_path) -> None:
    docx = _build_docx(
        tmp_path / "bullets.docx",
        [
            ("p", "Reasons:"),
            ("bullet", "First reason"),
            ("bullet", "Second reason"),
        ],
    )

    text = extract_document(str(docx))

    assert "• First reason" in text
    assert "• Second reason" in text


def test_docx_numbered_lists_are_numbered(tmp_path) -> None:
    docx = _build_docx(
        tmp_path / "steps.docx",
        [
            ("number", "Open the app"),
            ("number", "Ask a question"),
            ("p", "Interruption resets the counter."),
            ("number", "Back to a new list"),
        ],
    )

    text = extract_document(str(docx))

    assert "1. Open the app" in text
    assert "2. Ask a question" in text
    # The counter restarts after an intervening paragraph.
    assert "1. Back to a new list" in text


def test_docx_tables_are_rendered_as_text_tables(tmp_path) -> None:
    docx = _build_docx(
        tmp_path / "table.docx",
        [
            ("p", "Overview:"),
            (
                "table",
                [
                    ["Feature", "Status"],
                    ["PDF", "Supported"],
                    ["DOCX", "Supported"],
                ],
            ),
        ],
    )

    text = extract_document(str(docx))

    assert "Table:" in text
    assert "Feature | Status" in text
    assert "PDF | Supported" in text
    assert "DOCX | Supported" in text
    assert "Overview:" in text
    assert text.index("Overview:") < text.index("Table:")


def test_docx_logical_order_is_preserved(tmp_path) -> None:
    docx = _build_docx(
        tmp_path / "structure.docx",
        [
            ("h", (1, "Heading")),
            ("p", "Paragraph"),
            ("bullet", "Item 1"),
            ("bullet", "Item 2"),
            ("table", [["Column A", "Column B"], ["Value 1", "Value 2"]]),
        ],
    )

    text = extract_document(str(docx))

    assert text.index("# Heading") < text.index("Paragraph")
    assert text.index("Paragraph") < text.index("• Item 1")
    assert text.index("• Item 2") < text.index("Table:")
    assert text.index("Table:") < text.index("Value 1")


def test_docx_empty_document_extracts_to_empty(tmp_path) -> None:
    from docx import Document

    docx = tmp_path / "empty.docx"
    Document().save(str(docx))

    assert extract_document(str(docx)) == ""


def test_docx_unreadable_raises_document_error(tmp_path) -> None:
    docx = tmp_path / "bad.docx"
    docx.write_bytes(b"not a zip container")

    with pytest.raises(DocumentInputError):
        extract_document(str(docx))


# ── Unsupported formats ──────────────────────────────────────────────────────


def test_legacy_doc_is_recognised_but_unsupported(tmp_path) -> None:
    doc = tmp_path / "legacy.doc"
    doc.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")  # OLE compound magic

    with pytest.raises(UnsupportedDocumentError):
        extract_document(str(doc))


def test_unknown_extension_is_unsupported(tmp_path) -> None:
    unknown = tmp_path / "data.xyz"
    unknown.write_bytes(b"whatever")

    with pytest.raises(UnsupportedDocumentError):
        extract_document(str(unknown))


def test_is_document_path() -> None:
    assert is_document_path("notes.pdf")
    assert is_document_path("notes.docx")
    assert is_document_path("notes.doc")
    assert is_document_path("UPPER.PDF")
    assert not is_document_path("notes.txt")
    assert not is_document_path("photo.png")
    assert not is_document_path("no-extension")


def test_document_exts_cover_pdf_docx_doc() -> None:
    assert DOCUMENT_EXTS == {".pdf", ".docx", ".doc"}


# ── Large-document cap ───────────────────────────────────────────────────────


def test_small_document_passes_through_uncapped() -> None:
    text = "A short document."
    capped, truncated = cap_document_text(text)

    assert capped == text
    assert truncated is False


def test_large_document_is_truncated_at_block_boundary(tmp_path) -> None:
    pdf = tmp_path / "big.pdf"
    # One ~3000-char block per page: 40 pages ≈ 120k chars, over the budget.
    page = "Word. " * 500
    pdf.write_bytes(make_pdf([page] * 40))

    text = extract_pdf_text(str(pdf))
    capped, truncated = cap_document_text(text)

    assert truncated is True
    assert len(capped) <= MAX_DOCUMENT_CHARS
    # Whole pages are kept: the first page is present and not split.
    assert "[Page 1]" in capped
    first_page_end = capped.index("[Page 2]") if "[Page 2]" in capped else len(capped)
    first_page = capped[len("[Page 1]\n") : first_page_end]
    assert first_page.count("Word.") == 500


def test_cap_never_exceeds_budget_even_for_one_giant_block() -> None:
    text = "x" * (MAX_DOCUMENT_CHARS * 3)
    capped, truncated = cap_document_text(text)

    assert truncated is True
    assert len(capped) == MAX_DOCUMENT_CHARS


# ── Prompt composition ───────────────────────────────────────────────────────


def test_document_prompt_separates_question_and_document() -> None:
    prompt = build_document_prompt(
        "What are the three main conclusions?",
        "Conclusion one.\nConclusion two.",
        "notes.pdf",
    )

    assert prompt.startswith("USER QUESTION:\nWhat are the three main conclusions?")
    assert 'ATTACHED DOCUMENT: "notes.pdf"' in prompt
    assert '"""\nConclusion one.\nConclusion two.\n"""' in prompt
    assert "using ONLY the document content above" in prompt
    assert "could not be found in the document" in prompt


def test_document_prompt_notes_truncation_when_capped() -> None:
    prompt = build_document_prompt(
        "Summarize.", "Partial content.", "big.pdf", truncated=True
    )

    assert "[The document was longer than could be included" in prompt


def test_document_prompt_omits_truncation_note_by_default() -> None:
    prompt = build_document_prompt("Summarize.", "Full content.", "notes.pdf")

    assert "[The document was longer" not in prompt
