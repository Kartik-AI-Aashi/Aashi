"""Verify the UI-02B chat window.

Creation, the three chrome regions, message stacking with correct alignment,
auto-scroll on new messages, the canvas backdrop, and re-styling when the theme
changes.
"""

from __future__ import annotations

import threading

import pytest

pytest.importorskip("pytestqt", reason="pytest-qt is required for GUI tests")

from PySide6.QtCore import Qt  # noqa: E402
from PySide6.QtGui import QImage  # noqa: E402
from PySide6.QtWidgets import QSizePolicy, QToolButton  # noqa: E402

from aashi.ui.components.chat_bubble import (  # noqa: E402
    CHAT_BUBBLE_MAX_WIDTH_PX,
    ChatBubble,
    ChatRole,
)
from aashi.ui.components.header import AashiHeader  # noqa: E402
from aashi.ui.components.input import AashiInput  # noqa: E402
from aashi.ui.services.image_input import ImageInputError  # noqa: E402
from aashi.ui.services.openrouter import OpenRouterError  # noqa: E402
from aashi.ui.themes.definitions import DARK_THEME, LIGHT_THEME  # noqa: E402
from aashi.ui.themes.manager import ThemeManager  # noqa: E402
from aashi.ui.themes.mode import ColorScheme, ThemeMode  # noqa: E402
from aashi.ui.themes.system import SystemPreferences  # noqa: E402
from aashi.ui.windows import chat_window  # noqa: E402
from aashi.ui.windows.chat_window import (  # noqa: E402
    AashiChatWindow,
    _DOCUMENT_READ_ERROR,
    _SCANNED_PDF_MESSAGE,
    _THINKING_PHRASES,
    _THINKING_PUNCTUATION,
    _UNSUPPORTED_DOCUMENT_FORMAT,
)

from pdf_fixtures import make_pdf  # noqa: E402


def _manager() -> ThemeManager:
    return ThemeManager(system_reader=lambda: SystemPreferences(
        color_scheme=ColorScheme.DARK,
        high_contrast=False,
        reduced_motion=False,
    ))


def _fail_without_network(_message: str) -> str:
    raise OpenRouterError("OpenRouter calls are disabled in tests")


@pytest.fixture(autouse=True)
def _no_real_network(monkeypatch) -> None:
    """Keep every chat-window test off the real OpenRouter API.

    The window answers on a worker thread, so any test that presses Enter
    would otherwise perform a live network call. The default stub fails,
    exercising the error-fallback path unless a test overrides
    ``chat_window.fetch_response_hybrid`` or ``chat_window.fetch_vision_response``
    with its own mock.
    """
    monkeypatch.setattr(chat_window, "fetch_response_hybrid", _fail_without_network)
    monkeypatch.setattr(chat_window, "fetch_vision_response", _fail_without_network)


def test_window_creates_correctly(qtbot) -> None:
    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)

    assert window.message_count == 0
    assert window.styleSheet()


def test_adding_messages_increases_the_count(qtbot) -> None:
    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)

    window.add_message("hello", is_user=True)
    assert window.message_count == 1

    window.add_message("hi there", is_user=False)
    window.add_message("third message", is_user=True)
    assert window.message_count == 3


def test_scroll_area_exists(qtbot) -> None:
    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)

    assert window.scroll_area is not None
    assert window.scroll_area.widget() is not None
    assert window.scroll_area.widgetResizable()


def test_input_exists(qtbot) -> None:
    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)

    assert isinstance(window.input, AashiInput)


def test_header_exists(qtbot) -> None:
    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)

    assert isinstance(window.header, AashiHeader)
    assert window.header.title == "Aashi"


def test_user_and_assistant_messages_align_to_their_sides(qtbot) -> None:
    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)

    user = window.add_message("from the user", is_user=True)
    assistant = window.add_message("from Aashi", is_user=False)

    assert user.alignment == Qt.AlignmentFlag.AlignRight
    assert assistant.alignment == Qt.AlignmentFlag.AlignLeft


def test_bubble_widths_follow_content_length(qtbot) -> None:
    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)
    window.resize(800, 500)
    window.show()

    short = window.add_message("hi", is_user=True)
    # Long enough that the label's natural wrap width exceeds the bubble cap in
    # the offscreen font metrics, so the bubble truly clamps to the cap.
    long = window.add_message("word " * 2000, is_user=False)
    qtbot.wait(10)

    # Neither bubble spans the viewport: the short message hugs its text while
    # the long one is capped at the reading width and wraps. The window must
    # not override the bubble's content-hugging size policy.
    assert short.sizePolicy().horizontalPolicy() == QSizePolicy.Policy.Maximum
    assert short.width() < 300
    assert long.width() == CHAT_BUBBLE_MAX_WIDTH_PX
    assert short.width() < long.width()


def test_window_backdrop_uses_the_canvas_token(qtbot) -> None:
    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)

    assert DARK_THEME.palette.surface.canvas.css in window.styleSheet()


def test_window_restyles_when_the_theme_changes(qtbot) -> None:
    manager = _manager()
    window = AashiChatWindow(manager)
    qtbot.addWidget(window)
    assert DARK_THEME.palette.surface.canvas.css in window.styleSheet()

    manager.set_mode(ThemeMode.LIGHT)

    assert LIGHT_THEME.palette.surface.canvas.css in window.styleSheet()


def test_new_message_scrolls_to_bottom(qtbot) -> None:
    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)
    window.resize(400, 240)
    window.show()

    for index in range(20):
        window.add_message(f"message {index}", is_user=index % 2 == 0)
    qtbot.wait(50)

    scrollbar = window.scroll_area.verticalScrollBar()
    assert scrollbar.value() == scrollbar.maximum()


def test_pressing_enter_sends_the_input_as_a_user_message(qtbot) -> None:
    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)
    window.input.setText("Hello Aashi")

    qtbot.keyClick(window.input, Qt.Key.Key_Return)

    assert window.message_count == 1  # simple message → no thinking bubble
    assert window.input.text() == ""

    message_layout = window.scroll_area.widget().layout()
    bubble = message_layout.itemAt(0).widget()
    assert isinstance(bubble, ChatBubble)
    assert bubble.role is ChatRole.USER
    assert bubble.text == "Hello Aashi"


def test_pressing_enter_ignores_an_empty_input(qtbot) -> None:
    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)

    qtbot.keyClick(window.input, Qt.Key.Key_Return)

    assert window.message_count == 0


def test_pressing_enter_ignores_whitespace_only_input(qtbot) -> None:
    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)
    window.input.setText("   ")

    qtbot.keyClick(window.input, Qt.Key.Key_Return)

    assert window.message_count == 0


def test_sending_a_message_gets_a_real_ai_reply(qtbot, monkeypatch) -> None:
    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)
    monkeypatch.setattr(
        chat_window, "fetch_response_hybrid", lambda message: f"echo: {message}"
    )
    window.input.setText("hello")

    qtbot.keyClick(window.input, Qt.Key.Key_Return)

    # Simple message → no thinking bubble; only the user message is present
    # right away. The reply arrives via a typewriter that creates one bubble.
    assert window.message_count == 1

    message_layout = window.scroll_area.widget().layout()
    # The typewriter creates one bubble and fills it progressively; both
    # conditions together only hold once it has finished, so this cannot
    # return early while the reply is still pending or mid-reveal.
    qtbot.waitUntil(
        lambda: window.message_count == 2 and window._typewriter_bubble is None,
        timeout=4000,
    )
    assistant = message_layout.itemAt(1).widget()
    assert isinstance(assistant, ChatBubble)
    assert assistant.role is ChatRole.ASSISTANT
    # The reply echoes the request, which now carries the remembered transcript.
    assert assistant.text == "echo: User: hello"
    assert window.message_count == 2  # user + typewriter bubble


def test_failed_ai_request_shows_the_error_fallback(qtbot) -> None:
    # The autouse _no_real_network fixture makes fetch_response_hybrid fail, so the
    # window must degrade to the fallback bubble instead of crashing.
    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)
    window.input.setText("hello")

    qtbot.keyClick(window.input, Qt.Key.Key_Return)

    # Simple message → no thinking bubble; error bubble is added directly.
    assert window.message_count == 1

    message_layout = window.scroll_area.widget().layout()
    qtbot.waitUntil(
        lambda: window.message_count == 2,
        timeout=1000,
    )
    fallback = message_layout.itemAt(1).widget()
    assert isinstance(fallback, ChatBubble)
    assert fallback.role is ChatRole.ASSISTANT
    assert fallback.text == "Something went wrong"


# ── Composer UI features ─────────────────────────────────────────────────────


def test_input_has_placeholder_text(qtbot) -> None:
    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)

    assert window.input.placeholderText() == "Message Aashi..."


def test_send_button_exists(qtbot) -> None:
    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)

    assert window._send_button is not None
    assert window._send_button.text() == "➤"


def test_send_button_sends_message(qtbot) -> None:
    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)
    window.input.setText("click send")

    window._send_button.click()

    assert window.message_count == 1  # simple message → no thinking bubble
    assert window.input.text() == ""

    message_layout = window.scroll_area.widget().layout()
    bubble = message_layout.itemAt(0).widget()
    assert isinstance(bubble, ChatBubble)
    assert bubble.text == "click send"


def test_attach_button_exists(qtbot) -> None:
    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)

    assert window._attach_button is not None
    assert window._attach_button.text() == "📎"


def test_attachment_preview_is_hidden_by_default(qtbot) -> None:
    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)

    assert window._attachment_preview.isHidden()
    assert window._selected_file is None


def test_attach_stores_selected_file_and_shows_preview(qtbot, monkeypatch) -> None:
    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)

    monkeypatch.setattr(
        "aashi.ui.windows.chat_window.QFileDialog.getOpenFileName",
        lambda *a, **k: ("/tmp/test.txt", "Text files"),
    )

    window._on_attach_clicked()

    assert window._selected_file == "/tmp/test.txt"
    assert not window._attachment_preview.isHidden()
    assert window._attachment_preview.filename == "test.txt"


def test_attach_cancelled_does_not_store_file(qtbot, monkeypatch) -> None:
    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)

    monkeypatch.setattr(
        "aashi.ui.windows.chat_window.QFileDialog.getOpenFileName",
        lambda *a, **k: ("", ""),
    )

    window._on_attach_clicked()

    assert window._selected_file is None
    assert window._attachment_preview.isHidden()


# ── File reading ─────────────────────────────────────────────────────────────


def test_read_file_content_text(tmp_path) -> None:
    from aashi.ui.windows.chat_window import AashiChatWindow

    window = AashiChatWindow(_manager())
    txt = tmp_path / "notes.txt"
    txt.write_text("hello world", encoding="utf-8")
    window._selected_file = str(txt)

    assert window._read_file_content() == "hello world"


def test_read_file_content_pdf(tmp_path) -> None:
    from unittest.mock import MagicMock

    from aashi.ui.windows.chat_window import AashiChatWindow

    window = AashiChatWindow(_manager())
    pdf = tmp_path / "doc.pdf"
    pdf.write_bytes(b"%PDF-1.4 dummy")
    window._selected_file = str(pdf)

    fake_page = MagicMock()
    fake_page.extract_text.return_value = " extracted text "
    fake_reader = MagicMock()
    fake_reader.pages = [fake_page]

    with __import__("unittest.mock", fromlist=["patch"]).patch(
        "pypdf.PdfReader", return_value=fake_reader
    ):
        result = window._read_file_content()

    assert result == "extracted text"


def test_read_file_content_unreadable(tmp_path) -> None:
    from aashi.ui.windows.chat_window import AashiChatWindow

    window = AashiChatWindow(_manager())
    bad = tmp_path / "bad.txt"
    bad.write_bytes(b"\x80\x81\x82")  # invalid UTF-8
    window._selected_file = str(bad)

    assert window._read_file_content() == "[Unable to read file content]"


def test_submit_message_includes_file_content(
    qtbot, monkeypatch, tmp_path
) -> None:
    from aashi.ui.windows.chat_window import AashiChatWindow

    monkeypatch.setattr(
        chat_window, "fetch_response_hybrid", lambda message: f"echo: {message}"
    )
    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)

    txt = tmp_path / "data.txt"
    txt.write_text("SECRET_DATA", encoding="utf-8")
    window._selected_file = str(txt)
    window._attachment_preview.set_attachment(str(txt))

    window.input.setText("summarize this")
    qtbot.keyClick(window.input, Qt.Key.Key_Return)

    # The user bubble shows the typed text plus a paperclip line naming the
    # file; file content itself is only in the AI message.
    user_bubble = window.scroll_area.widget().layout().itemAt(0).widget()
    assert user_bubble.text == "summarize this\n📎 data.txt"
    assert "SECRET_DATA" not in user_bubble.text

    # File is cleared and preview hidden.
    assert window._selected_file is None
    assert window._attachment_preview.isHidden()

    # The AI mock received the file content.
    layout = window.scroll_area.widget().layout()
    qtbot.waitUntil(
        lambda: window.message_count == 2 and window._typewriter_bubble is None,
        timeout=4000,
    )
    assistant = layout.itemAt(1).widget()
    assert "SECRET_DATA" in assistant.text


def test_submit_without_file_skips_content(qtbot, monkeypatch) -> None:
    from aashi.ui.windows.chat_window import AashiChatWindow

    monkeypatch.setattr(
        chat_window, "fetch_response_hybrid", lambda message: f"echo: {message}"
    )
    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)

    window.input.setText("hello")
    qtbot.keyClick(window.input, Qt.Key.Key_Return)

    layout = window.scroll_area.widget().layout()
    qtbot.waitUntil(
        lambda: window.message_count == 2 and window._typewriter_bubble is None,
        timeout=4000,
    )
    assistant = layout.itemAt(1).widget()
    assert "[Attached File Content]" not in assistant.text
    assert assistant.text == "echo: User: hello"


# ── Attachment preview + vision ──────────────────────────────────────────────


def _make_png(path) -> None:
    """Write a small valid PNG for preview/vision tests."""
    image = QImage(120, 60, QImage.Format.Format_RGB32)
    image.fill(Qt.GlobalColor.red)
    assert image.save(str(path), "PNG")


def test_image_attach_shows_thumbnail_and_filename(
    qtbot, monkeypatch, tmp_path
) -> None:
    from aashi.ui.windows.chat_window import AashiChatWindow

    png = tmp_path / "photo.png"
    _make_png(png)
    monkeypatch.setattr(
        "aashi.ui.windows.chat_window.QFileDialog.getOpenFileName",
        lambda *a, **k: (str(png), "Images (*.png)"),
    )

    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)
    window._on_attach_clicked()

    assert window._selected_file == str(png)
    assert not window._attachment_preview.isHidden()
    assert window._attachment_preview.filename == "photo.png"
    assert not window._attachment_preview._thumb.pixmap().isNull()


def test_remove_attachment_keeps_typed_text(qtbot, tmp_path) -> None:
    from aashi.ui.windows.chat_window import AashiChatWindow

    png = tmp_path / "photo.png"
    _make_png(png)
    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)

    window._selected_file = str(png)
    window._attachment_preview.set_attachment(str(png))
    window.input.setText("keep this text")
    assert not window._attachment_preview.isHidden()

    close = window._attachment_preview.findChild(QToolButton)
    close.click()

    assert window._selected_file is None
    assert window._attachment_preview.isHidden()
    assert window.input.text() == "keep this text"
    assert window.message_count == 0


def test_submit_image_routes_to_vision(qtbot, monkeypatch, tmp_path) -> None:
    from aashi.ui.windows.chat_window import AashiChatWindow

    png = tmp_path / "photo.png"
    _make_png(png)
    seen: dict[str, object] = {}

    def fake_encode(_path: str) -> str:
        return "data:image/png;base64,QUFB"

    def fake_vision(message: str, data_uri: str) -> str:
        seen["data_uri"] = data_uri
        return f"seen: {message}"

    monkeypatch.setattr(chat_window, "encode_image_data_uri", fake_encode)
    monkeypatch.setattr(chat_window, "fetch_vision_response", fake_vision)

    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)
    window._selected_file = str(png)
    window._attachment_preview.set_attachment(str(png))
    window.input.setText("What is in this image?")
    qtbot.keyClick(window.input, Qt.Key.Key_Return)

    layout = window.scroll_area.widget().layout()
    user_bubble = layout.itemAt(0).widget()
    assert user_bubble.text == "What is in this image?\n🖼 photo.png"

    # The attachment is consumed: preview clears, state resets.
    assert window._selected_file is None
    assert window._attachment_preview.isHidden()

    qtbot.waitUntil(
        lambda: window._thinking_bubble is None
        and window._typewriter_bubble is None
        and window.message_count == 2,
        timeout=4000,
    )
    assert seen.get("data_uri") == "data:image/png;base64,QUFB"
    assistant = layout.itemAt(1).widget()
    assert "What is in this image?" in assistant.text


def test_image_decode_failure_shows_friendly_error(
    qtbot, monkeypatch, tmp_path
) -> None:
    from aashi.ui.windows.chat_window import AashiChatWindow

    png = tmp_path / "broken.png"
    png.write_bytes(b"\xff\xd8\xff\xe0 not a real image")
    vision_called: list[str] = []

    def fail_encode(_path: str) -> str:
        raise ImageInputError("cannot decode")

    monkeypatch.setattr(chat_window, "encode_image_data_uri", fail_encode)
    monkeypatch.setattr(
        chat_window,
        "fetch_vision_response",
        lambda msg, uri: vision_called.append(uri) or "unreachable",
    )

    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)
    window._selected_file = str(png)
    window._attachment_preview.set_attachment(str(png))
    window.input.setText("What is in this image?")
    qtbot.keyClick(window.input, Qt.Key.Key_Return)

    qtbot.waitUntil(
        lambda: window._thinking_bubble is None
        and window._typewriter_bubble is None
        and window.message_count == 2,
        timeout=4000,
    )
    layout = window.scroll_area.widget().layout()
    assistant = layout.itemAt(1).widget()
    assert assistant.text == "I couldn't read that image."
    assert vision_called == []
    # The attachment is still consumed so the composer resets cleanly.
    assert window._selected_file is None
    assert window._attachment_preview.isHidden()


# ── Document understanding ───────────────────────────────────────────────────


def _send_document(window, qtbot, path, question) -> None:
    """Stage ``path``, type ``question`` and press Enter."""
    qtbot.addWidget(window)  # let qtbot own the window for clean teardown
    window._selected_file = str(path)
    window._attachment_preview.set_attachment(str(path))
    window.input.setText(question)
    qtbot.keyClick(window.input, Qt.Key.Key_Return)


def _wait_for_reply(qtbot, window) -> None:
    qtbot.waitUntil(
        lambda: window._thinking_bubble is None
        and window._typewriter_bubble is None
        and window.message_count == 2,
        timeout=4000,
    )


def test_submit_pdf_routes_extracted_content_to_ai(
    qtbot, monkeypatch, tmp_path
) -> None:
    from aashi.ui.windows.chat_window import AashiChatWindow

    pdf = tmp_path / "notes.pdf"
    pdf.write_bytes(
        make_pdf(
            [
                "Photosynthesis converts sunlight into chemical energy.",
                "Chlorophyll absorbs red and blue light.",
            ]
        )
    )
    seen: dict[str, str] = {}

    def fake_hybrid(message: str) -> str:
        seen["message"] = message
        return "Here is the summary."

    monkeypatch.setattr(chat_window, "fetch_response_hybrid", fake_hybrid)

    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)
    _send_document(window, qtbot, pdf, "Summarize this document.")

    # The attachment is consumed: preview clears, state resets.
    assert window._selected_file is None
    assert window._attachment_preview.isHidden()

    _wait_for_reply(qtbot, window)

    message = seen["message"]
    # The AI received the question, both pages of extracted content, and the
    # file name — never the raw filesystem path.
    assert "Summarize this document." in message
    assert "Photosynthesis converts sunlight" in message
    assert "Chlorophyll absorbs red and blue light" in message
    assert "[Page 1]" in message
    assert "[Page 2]" in message
    assert "notes.pdf" in message
    assert str(tmp_path) not in message
    assert "USER QUESTION:" in message
    assert 'ATTACHED DOCUMENT: "notes.pdf"' in message

    layout = window.scroll_area.widget().layout()
    assistant = layout.itemAt(1).widget()
    assert "Here is the summary." in assistant.text


def test_submit_docx_sends_extracted_structure_to_ai(
    qtbot, monkeypatch, tmp_path
) -> None:
    from aashi.ui.windows.chat_window import AashiChatWindow
    from docx import Document

    docx = tmp_path / "assignment.docx"
    document = Document()
    document.add_heading("Main Topics", level=1)
    document.add_paragraph("Biology covers living organisms.")
    document.add_paragraph("Photosynthesis", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Topic"
    table.cell(0, 1).text = "Chapter"
    table.cell(1, 0).text = "Cells"
    table.cell(1, 1).text = "Chapter 2"
    document.save(str(docx))

    seen: dict[str, str] = {}
    monkeypatch.setattr(
        chat_window,
        "fetch_response_hybrid",
        lambda message: seen.__setitem__("message", message) or "summary",
    )

    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)
    _send_document(window, qtbot, docx, "What are the main points?")
    _wait_for_reply(qtbot, window)

    message = seen["message"]
    # The DOCX structure — heading, paragraph, list item, table — reaches the
    # AI in document order, without the raw path.
    assert "What are the main points?" in message
    assert "# Main Topics" in message
    assert "Biology covers living organisms." in message
    assert "• Photosynthesis" in message
    assert "Topic | Chapter" in message
    assert "Cells | Chapter 2" in message
    assert str(tmp_path) not in message


def test_submit_unreadable_document_shows_friendly_error(
    qtbot, tmp_path
) -> None:
    from aashi.ui.windows.chat_window import AashiChatWindow

    pdf = tmp_path / "broken.pdf"
    pdf.write_bytes(b"this is not a pdf")

    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)
    _send_document(window, qtbot, pdf, "Summarize this document.")
    _wait_for_reply(qtbot, window)

    layout = window.scroll_area.widget().layout()
    assistant = layout.itemAt(1).widget()
    assert assistant.text == _DOCUMENT_READ_ERROR
    assert window._selected_file is None
    assert window._attachment_preview.isHidden()


def test_submit_scanned_pdf_shows_explanation(qtbot, tmp_path) -> None:
    from aashi.ui.windows.chat_window import AashiChatWindow

    scan = tmp_path / "scan.pdf"
    scan.write_bytes(make_pdf(["", ""]))  # image-based pages, no text

    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)
    _send_document(window, qtbot, scan, "What does this show?")
    _wait_for_reply(qtbot, window)

    layout = window.scroll_area.widget().layout()
    assistant = layout.itemAt(1).widget()
    assert assistant.text == _SCANNED_PDF_MESSAGE


def test_submit_unsupported_doc_shows_message(qtbot, tmp_path) -> None:
    from aashi.ui.windows.chat_window import AashiChatWindow

    doc = tmp_path / "legacy.doc"
    doc.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")  # OLE magic bytes

    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)
    _send_document(window, qtbot, doc, "Read this document.")
    _wait_for_reply(qtbot, window)

    layout = window.scroll_area.widget().layout()
    assistant = layout.itemAt(1).widget()
    assert assistant.text == _UNSUPPORTED_DOCUMENT_FORMAT
    assert window._selected_file is None


# ── Thinking indicator ───────────────────────────────────────────────────────


_ALL_THINKING_TEXTS = frozenset(
    f"{phrase}{ending}"
    for phrase in _THINKING_PHRASES
    for ending in _THINKING_PUNCTUATION
)


def _gate_the_reply(monkeypatch, prefix: str = "echo: ") -> threading.Event:
    """Block the worker thread's reply until the returned event is set."""
    gate = threading.Event()
    monkeypatch.setattr(
        chat_window,
        "fetch_response_hybrid",
        lambda message: (gate.wait(5) and None) or f"{prefix}{message}",
    )
    return gate


def test_sending_shows_a_thinking_bubble_while_the_reply_is_pending(
    qtbot, monkeypatch
) -> None:
    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)
    gate = _gate_the_reply(monkeypatch)
    window.input.setText("explain this complex topic in detail")

    qtbot.keyClick(window.input, Qt.Key.Key_Return)

    # The reply is still blocked, so the placeholder bubble is on screen with
    # one of the randomly chosen phrases.
    assert window.message_count == 2  # user + thinking bubble
    message_layout = window.scroll_area.widget().layout()
    thinking = message_layout.itemAt(1).widget()
    assert isinstance(thinking, ChatBubble)
    assert thinking.role is ChatRole.ASSISTANT
    assert thinking.text in _ALL_THINKING_TEXTS

    gate.set()
    qtbot.waitUntil(
        lambda: window._thinking_bubble is None,
        timeout=1000,
    )


def test_thinking_bubble_removed_then_typewriter_creates_reply(
    qtbot, monkeypatch
) -> None:
    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)
    gate = _gate_the_reply(monkeypatch)
    window.input.setText("explain this complex topic in detail")
    qtbot.keyClick(window.input, Qt.Key.Key_Return)

    message_layout = window.scroll_area.widget().layout()
    placeholder = message_layout.itemAt(1).widget()
    assert isinstance(placeholder, ChatBubble)
    assert placeholder.text in _ALL_THINKING_TEXTS

    gate.set()
    qtbot.waitUntil(
        lambda: window._typewriter_bubble is None and window._thinking_bubble is None,
        timeout=4000,
    )

    # Thinking bubble was removed; reply is in a new typewriter bubble.
    assistant = message_layout.itemAt(1).widget()
    assert assistant is not placeholder
    assert assistant.role is ChatRole.ASSISTANT
    assert window.message_count == 2  # user + typewriter bubble


def test_thinking_phrase_composes_from_list_and_punctuation(qtbot, monkeypatch) -> None:
    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)

    # Deterministic choice: first phrase, first ending. Verifies the phrase is
    # actually composed at send time from the sanctioned lists.
    monkeypatch.setattr(chat_window.random, "choice", lambda seq: seq[0])

    assert (
        chat_window._random_thinking_phrase()
        == f"{_THINKING_PHRASES[0]}{_THINKING_PUNCTUATION[0]}"
    )


def test_thinking_lists_offer_variety() -> None:
    assert len(_THINKING_PHRASES) >= 3
    assert len(set(_THINKING_PUNCTUATION)) >= 2


# ── Reply formatting (Fix 3) ──────────────────────────────────────────────────


def test_format_assistant_reply_breaks_sentences_into_paragraphs() -> None:
    from aashi.ui.windows.chat_window import _format_assistant_reply

    formatted = _format_assistant_reply(
        "First sentence. Second one? And then! Finally done."
    )
    assert (
        formatted
        == "First sentence.\n\nSecond one?\n\nAnd then!\n\nFinally done."
    )


def test_format_assistant_reply_keeps_numbered_items_intact() -> None:
    from aashi.ui.windows.chat_window import _format_assistant_reply

    formatted = _format_assistant_reply(
        "Here is a list: 1. first 2. second and 3. third."
    )
    # Markers survive as items and are never split by the sentence pass.
    assert "1. first" in formatted
    assert "2. second" in formatted
    assert "3. third." in formatted
    assert "1.\n\nfirst" not in formatted


def test_format_assistant_reply_collapses_excess_blank_lines() -> None:
    from aashi.ui.windows.chat_window import _format_assistant_reply

    formatted = _format_assistant_reply("Para one.\n\n\n\n\nPara two.")
    assert formatted == "Para one.\n\nPara two."


def test_format_assistant_reply_never_changes_the_wording() -> None:
    from aashi.ui.windows.chat_window import _format_assistant_reply

    source = "Exactly the same words here."
    formatted = _format_assistant_reply(source)
    # Only whitespace may change: every word and period is still present.
    assert formatted.replace("\n", "") == source


# ── Typewriter speed (Fix 1) ──────────────────────────────────────────────────


def test_typewriter_reveals_one_character_every_30ms() -> None:
    from aashi.ui.windows.chat_window import (
        _TYPEWRITER_CHARS_PER_TICK,
        _TYPEWRITER_INTERVAL_MS,
    )

    # 25–40 ms per character: one character per tick.
    assert _TYPEWRITER_CHARS_PER_TICK == 1
    assert 25 <= _TYPEWRITER_INTERVAL_MS <= 40


def test_typewriter_pauses_after_punctuation(qtbot) -> None:
    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)

    window._start_typewriter("a,b")
    # Tick 1 reveals "a"; the comma is revealed next and triggers a pause.
    window._typewriter_tick()
    assert window._typewriter_bubble.text == "a"
    window._typewriter_tick()
    assert window._typewriter_pause > 0
    # A pause tick holds without revealing more text.
    window._typewriter_tick()
    assert window._typewriter_bubble.text == "a,"


# ── Conversation memory (Fix 4) ───────────────────────────────────────────────


def test_conversation_history_feeds_the_next_request(qtbot, monkeypatch) -> None:
    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)
    sent: list[str] = []
    monkeypatch.setattr(
        chat_window,
        "fetch_response_hybrid",
        lambda message: sent.append(message) or "reply",
    )

    window.input.setText("first question")
    qtbot.keyClick(window.input, Qt.Key.Key_Return)
    qtbot.waitUntil(
        lambda: window.message_count == 2 and window._typewriter_bubble is None,
        timeout=4000,
    )
    # The user turn and the assistant reply are both remembered.
    assert [turn["role"] for turn in window._chat_history] == ["user", "assistant"]

    window.input.setText("second question")
    qtbot.keyClick(window.input, Qt.Key.Key_Return)
    qtbot.waitUntil(
        lambda: window._typewriter_bubble is None and len(sent) == 2,
        timeout=4000,
    )

    # The second request replays both turns, newest last.
    assert "first question" in sent[1]
    assert "second question" in sent[1]
    assert sent[1].index("first question") < sent[1].index("second question")


# ── Code-block rendering (code fences preserved) ───────────────────────────


def test_code_blocks_preserve_whitespace_in_reply(qtbot, monkeypatch) -> None:
    """Fenced code blocks in an AI reply survive the renderer with whitespace intact.

    The ``<pre>`` conversion must pass through the full path — AI reply →
    ``_format_assistant_reply`` → ``render_markdown`` → typewriter →
    ``QLabel`` ``setText`` — so the bubble shows the code block aligned.
    """
    flowchart_reply = (
        "Here is the process:\n\n"
        "```text\n"
        "        Start\n"
        "          |\n"
        "          v\n"
        "     Photosynthesis\n"
        "```\n\n"
        "Done."
    )
    monkeypatch.setattr(
        chat_window, "fetch_response_hybrid", lambda _message: flowchart_reply
    )
    window = AashiChatWindow(_manager())
    qtbot.addWidget(window)
    window.input.setText("give me a flowchart")

    qtbot.keyClick(window.input, Qt.Key.Key_Return)

    qtbot.waitUntil(
        lambda: window._typewriter_bubble is None and window._thinking_bubble is None,
        timeout=4000,
    )
    message_layout = window.scroll_area.widget().layout()
    assistant = message_layout.itemAt(1).widget()
    assert isinstance(assistant, ChatBubble)
    # The raw label text contains the <pre> HTML — the renderer ran.
    assert "<pre>" in assistant.text
    assert "</pre>" in assistant.text
    # The code block content with alignment survives.
    assert "        Start" in assistant.text
    assert "          |" in assistant.text
    assert "          v" in assistant.text
    assert "     Photosynthesis" in assistant.text
